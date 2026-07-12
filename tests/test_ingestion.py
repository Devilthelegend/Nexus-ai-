"""Tests for the document ingestion pipeline and API."""

import pytest
from httpx import AsyncClient

from app.ai.vectorstore.factory import _memory_store

_AUTH = "/api/v1/auth"
_WS = "/api/v1/workspaces"
_PASSWORD = "s3cret-password"


@pytest.fixture(autouse=True)
def isolate_ingestion(tmp_path, monkeypatch):
    """Point uploads at a temp dir and reset the in-memory vector store."""
    from app.core.config import get_settings

    monkeypatch.setattr(get_settings(), "upload_dir", str(tmp_path))
    _memory_store.cache_clear()
    yield
    _memory_store.cache_clear()


async def _workspace(client: AsyncClient, email: str) -> dict[str, object]:
    """Register/login a user and create a workspace; return ids and headers."""
    await client.post(
        f"{_AUTH}/register", json={"email": email, "password": _PASSWORD}
    )
    login = await client.post(
        f"{_AUTH}/login", json={"email": email, "password": _PASSWORD}
    )
    headers = {"Authorization": f"Bearer {login.json()['access_token']}"}
    created = await client.post(_WS, json={"name": "KB"}, headers=headers)
    return {"workspace_id": created.json()["id"], "headers": headers}


def _docs_url(workspace_id: str) -> str:
    return f"{_WS}/{workspace_id}/documents"


async def test_upload_indexes_document(client: AsyncClient) -> None:
    ws = await _workspace(client, "ingest-owner@example.com")
    files = {"file": ("notes.txt", b"NexusAI ingestion pipeline test.", "text/plain")}

    response = await client.post(
        _docs_url(ws["workspace_id"]), files=files, headers=ws["headers"]
    )
    assert response.status_code == 201, response.text
    body = response.json()
    assert body["status"] == "indexed"
    assert body["chunk_count"] >= 1
    assert body["error"] is None


async def test_upload_empty_file_rejected(client: AsyncClient) -> None:
    ws = await _workspace(client, "empty-owner@example.com")
    files = {"file": ("empty.txt", b"", "text/plain")}

    response = await client.post(
        _docs_url(ws["workspace_id"]), files=files, headers=ws["headers"]
    )
    assert response.status_code == 400


async def test_upload_is_idempotent(client: AsyncClient) -> None:
    ws = await _workspace(client, "idem-owner@example.com")
    files = {"file": ("dup.txt", b"identical content", "text/plain")}

    first = await client.post(
        _docs_url(ws["workspace_id"]), files=files, headers=ws["headers"]
    )
    second = await client.post(
        _docs_url(ws["workspace_id"]),
        files={"file": ("dup.txt", b"identical content", "text/plain")},
        headers=ws["headers"],
    )
    assert first.json()["id"] == second.json()["id"]

    listing = await client.get(
        _docs_url(ws["workspace_id"]), headers=ws["headers"]
    )
    assert len(listing.json()) == 1


async def test_status_and_get(client: AsyncClient) -> None:
    ws = await _workspace(client, "status-owner@example.com")
    files = {"file": ("s.txt", b"content for status", "text/plain")}
    doc_id = (
        await client.post(
            _docs_url(ws["workspace_id"]), files=files, headers=ws["headers"]
        )
    ).json()["id"]

    got = await client.get(
        f"{_docs_url(ws['workspace_id'])}/{doc_id}", headers=ws["headers"]
    )
    assert got.status_code == 200
    status = await client.get(
        f"{_docs_url(ws['workspace_id'])}/{doc_id}/status", headers=ws["headers"]
    )
    assert status.status_code == 200
    assert status.json()["status"] == "indexed"


async def test_delete_document(client: AsyncClient) -> None:
    ws = await _workspace(client, "del-owner@example.com")
    files = {"file": ("d.txt", b"to be deleted", "text/plain")}
    doc_id = (
        await client.post(
            _docs_url(ws["workspace_id"]), files=files, headers=ws["headers"]
        )
    ).json()["id"]

    deleted = await client.delete(
        f"{_docs_url(ws['workspace_id'])}/{doc_id}", headers=ws["headers"]
    )
    assert deleted.status_code == 204
    missing = await client.get(
        f"{_docs_url(ws['workspace_id'])}/{doc_id}", headers=ws["headers"]
    )
    assert missing.status_code == 404


async def test_tenant_isolation_on_upload(client: AsyncClient) -> None:
    ws = await _workspace(client, "iso-a@example.com")
    outsider = await _workspace(client, "iso-b@example.com")
    files = {"file": ("x.txt", b"secret", "text/plain")}

    response = await client.post(
        _docs_url(ws["workspace_id"]), files=files, headers=outsider["headers"]
    )
    assert response.status_code == 404


def _make_pdf(text: str) -> bytes:
    """Build a minimal single-page PDF with extractable text (no extra deps)."""
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
    ]
    stream = b"BT /F1 24 Tf 72 720 Td (" + text.encode("latin-1") + b") Tj ET"
    objects.append(
        b"<< /Length "
        + str(len(stream)).encode()
        + b" >>\nstream\n"
        + stream
        + b"\nendstream"
    )
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    pdf = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf += str(index).encode() + b" 0 obj\n" + obj + b"\nendobj\n"
    xref_pos = len(pdf)
    size = len(objects) + 1
    pdf += b"xref\n0 " + str(size).encode() + b"\n0000000000 65535 f \n"
    for off in offsets:
        pdf += ("%010d 00000 n \n" % off).encode()
    pdf += (
        b"trailer\n<< /Size "
        + str(size).encode()
        + b" /Root 1 0 R >>\nstartxref\n"
        + str(xref_pos).encode()
        + b"\n%%EOF"
    )
    return bytes(pdf)


def _make_docx(text: str) -> bytes:
    """Build a minimal .docx document in memory."""
    import io

    from docx import Document as DocxDocument

    buffer = io.BytesIO()
    document = DocxDocument()
    document.add_paragraph(text)
    document.save(buffer)
    return buffer.getvalue()


async def test_upload_pdf_indexes(client: AsyncClient) -> None:
    ws = await _workspace(client, "pdf-owner@example.com")
    files = {
        "file": (
            "doc.pdf",
            _make_pdf("NexusAI pdf ingestion works end to end"),
            "application/pdf",
        )
    }

    response = await client.post(
        _docs_url(ws["workspace_id"]), files=files, headers=ws["headers"]
    )
    assert response.status_code == 201, response.text
    body = response.json()
    assert body["status"] == "indexed", body
    assert body["chunk_count"] >= 1
    assert body["error"] is None


async def test_upload_docx_indexes(client: AsyncClient) -> None:
    ws = await _workspace(client, "docx-owner@example.com")
    files = {
        "file": (
            "doc.docx",
            _make_docx("NexusAI docx ingestion works with hybrid retrieval"),
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document",
        )
    }

    response = await client.post(
        _docs_url(ws["workspace_id"]), files=files, headers=ws["headers"]
    )
    assert response.status_code == 201, response.text
    body = response.json()
    assert body["status"] == "indexed", body
    assert body["chunk_count"] >= 1
    assert body["error"] is None


async def test_unsupported_type_fails_and_reprocess(client: AsyncClient) -> None:
    ws = await _workspace(client, "fail-owner@example.com")
    files = {"file": ("data.bin", b"\x00\x01binary", "application/octet-stream")}

    created = await client.post(
        _docs_url(ws["workspace_id"]), files=files, headers=ws["headers"]
    )
    assert created.json()["status"] == "failed"
    assert created.json()["error"]

    doc_id = created.json()["id"]
    replay = await client.post(
        f"{_docs_url(ws['workspace_id'])}/{doc_id}/reprocess",
        headers=ws["headers"],
    )
    assert replay.status_code == 200
    assert replay.json()["status"] == "failed"
